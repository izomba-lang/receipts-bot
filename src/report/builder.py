from __future__ import annotations

import calendar
import logging
from datetime import date
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from src.report.fx import FxRateProvider
from src.report.recalc import recalc_xlsx

logger = logging.getLogger(__name__)

HEADER_FILL = PatternFill(start_color="305496", end_color="305496", fill_type="solid")
HEADER_FONT = Font(name="Arial", bold=True, size=11, color="FFFFFF")
BODY_FONT = Font(name="Arial", size=10)
TOTAL_FILL = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
TOTAL_FONT = Font(name="Arial", bold=True, size=10)
SUBTOTAL_FONT = Font(name="Arial", italic=True, size=10)
THIN_BORDER = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)
COL_WIDTHS = [4, 13, 20, 60, 14, 10, 18, 16]
HEADERS = [
    "#", "Date", "Provider", "Description",
    "Amount (orig.)", "Currency", "FX rate (1 unit → AED)", "Amount (AED)",
]


async def compute_aed_total(
    receipts: list[dict[str, Any]], fx_provider: FxRateProvider
) -> float:
    """Sum all receipts converted to AED using per-date mid-market rates —
    matches the report's AED TOTAL."""
    total = Decimal("0")
    for r in receipts:
        d = r["date"]
        if isinstance(d, str):
            d = date.fromisoformat(d)
        currency = (r.get("currency") or "AED").upper()
        amount = Decimal(str(r["amount"]))
        rate = await fx_provider.get_rate(d, currency)
        total += amount * rate
    return float(round(total, 2))


def _build_description(row: dict[str, Any]) -> str:
    cat = (row.get("category") or "other").capitalize()
    time_str = row.get("time") or ""
    if time_str and len(time_str) > 5:
        time_str = time_str[:5]
    from_loc = row.get("from_location")
    to_loc = row.get("to_location")

    parts = [cat]
    if time_str:
        parts.append(time_str)
    desc = " ".join(parts)

    if from_loc and to_loc:
        desc += f" — {from_loc} → {to_loc}"
    elif from_loc:
        desc += f" — {from_loc}"
    elif to_loc:
        desc += f" — {to_loc}"

    return desc


def _title_for_range(start: date, end: date) -> str:
    if start.day == 1 and end.month == start.month:
        last_day = calendar.monthrange(start.year, start.month)[1]
        if end.day == last_day:
            return f"{calendar.month_name[start.month]} {start.year}"
    return f"{start.isoformat()} – {end.isoformat()}"


async def build_report(
    receipts: list[dict[str, Any]],
    start: date,
    end: date,
    output_path: Path,
    fx_provider: FxRateProvider,
    title_override: str | None = None,
) -> Path:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Expense Report"

    title_label = title_override or _title_for_range(start, end)
    title_text = f"Expense Report — {title_label}"
    subtitle_text = (
        "Reimbursement currency: AED. "
        "FX rates applied as of payment date (mid-market)."
    )

    ws.merge_cells("A1:H1")
    title_cell = ws["A1"]
    title_cell.value = title_text
    title_cell.font = Font(name="Arial", bold=True, size=14)
    title_cell.alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("A2:H2")
    sub_cell = ws["A2"]
    sub_cell.value = subtitle_text
    sub_cell.font = Font(name="Arial", italic=True, size=9, color="808080")
    sub_cell.alignment = Alignment(horizontal="center", vertical="center")

    header_row = 4
    for col_idx, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER

    for col_idx, width in enumerate(COL_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    first_data_row = header_row + 1
    currencies_seen: set[str] = set()
    payment_methods: set[str] = set()

    for idx, receipt in enumerate(receipts):
        row_num = first_data_row + idx
        ws.row_dimensions[row_num].height = 30

        receipt_date = receipt["date"]
        if isinstance(receipt_date, str):
            receipt_date = date.fromisoformat(receipt_date)
        currency = (receipt.get("currency") or "AED").upper()
        amount = Decimal(str(receipt["amount"]))

        fx_rate = await fx_provider.get_rate(receipt_date, currency)

        currencies_seen.add(currency)
        pm = receipt.get("payment_method")
        if pm:
            payment_methods.add(pm)

        values: list[Any] = [
            idx + 1,
            receipt_date.strftime("%Y-%m-%d"),
            receipt.get("provider") or "",
            _build_description(receipt),
            float(amount),
            currency,
            float(fx_rate),
            None,
        ]

        for col_idx, val in enumerate(values, start=1):
            cell = ws.cell(row=row_num, column=col_idx, value=val)
            cell.font = BODY_FONT
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="center")

        ws.cell(row=row_num, column=5).number_format = "#,##0.00"
        ws.cell(row=row_num, column=7).number_format = "0.00000"

        aed_cell = ws.cell(row=row_num, column=8)
        aed_cell.value = f"=E{row_num}*G{row_num}"
        aed_cell.number_format = "#,##0.00"

    last_data_row = first_data_row + len(receipts) - 1

    total_row = last_data_row + 1
    ws.row_dimensions[total_row].height = 30
    for col_idx in range(1, 9):
        cell = ws.cell(row=total_row, column=col_idx)
        cell.fill = TOTAL_FILL
        cell.font = TOTAL_FONT
        cell.border = THIN_BORDER

    ws.cell(row=total_row, column=4, value="TOTAL").font = TOTAL_FONT
    total_cell = ws.cell(row=total_row, column=8)
    total_cell.value = f"=SUM(H{first_data_row}:H{last_data_row})"
    total_cell.number_format = '#,##0.00 "AED"'
    total_cell.font = TOTAL_FONT
    total_cell.fill = TOTAL_FILL

    subtotal_row = total_row + 1
    for ccy in sorted(currencies_seen):
        ws.row_dimensions[subtotal_row].height = 25
        label_cell = ws.cell(row=subtotal_row, column=4, value=f"Subtotal ({ccy})")
        label_cell.font = SUBTOTAL_FONT
        sub_cell = ws.cell(row=subtotal_row, column=8)
        sub_cell.value = (
            f'=SUMIF(F{first_data_row}:F{last_data_row},"{ccy}",'
            f"H{first_data_row}:H{last_data_row})"
        )
        sub_cell.number_format = "#,##0.00"
        sub_cell.font = SUBTOTAL_FONT
        subtotal_row += 1

    notes_row = subtotal_row + 2
    notes_header = ws.cell(row=notes_row, column=1, value="Notes:")
    notes_header.font = Font(name="Arial", bold=True, size=10)

    note_font = Font(name="Arial", italic=True, size=9, color="808080")
    note_row = notes_row + 1

    if payment_methods:
        cards = ", ".join(sorted(payment_methods))
        ws.merge_cells(f"A{note_row}:H{note_row}")
        ws.cell(row=note_row, column=1, value=f"• Payment cards: {cards}").font = note_font
        note_row += 1

    ws.merge_cells(f"A{note_row}:H{note_row}")
    ws.cell(
        row=note_row, column=1,
        value="• Mid-market rates from open currency-api on each payment date.",
    ).font = note_font
    note_row += 1

    dates_count: dict[str, int] = {}
    missing_routes = 0
    for r in receipts:
        d = str(r.get("date", ""))
        dates_count[d] = dates_count.get(d, 0) + 1
        if r.get("category") == "taxi" and not r.get("from_location") and not r.get("to_location"):
            missing_routes += 1

    anomalies: list[str] = []
    multi_days = [d for d, c in dates_count.items() if c > 2]
    if multi_days:
        anomalies.append(f"Multiple trips on: {', '.join(sorted(multi_days))}")
    if missing_routes:
        anomalies.append(f"{missing_routes} taxi receipt(s) without route information")

    for a in anomalies:
        ws.merge_cells(f"A{note_row}:H{note_row}")
        ws.cell(row=note_row, column=1, value=f"• {a}").font = note_font
        note_row += 1

    wb.save(output_path)
    recalc_xlsx(output_path)

    logger.info("Report saved to %s (%d receipts)", output_path, len(receipts))
    return output_path
