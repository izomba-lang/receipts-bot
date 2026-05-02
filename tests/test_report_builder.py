from __future__ import annotations

from datetime import date
from decimal import Decimal
from pathlib import Path
from unittest.mock import AsyncMock

import pytest
from openpyxl import load_workbook

from src.report.builder import build_report
from src.report.fx import FxRateProvider

SAMPLE_RECEIPTS = [
    {
        "id": 1,
        "date": "2026-04-20",
        "time": "08:15",
        "provider": "GetTaxi",
        "category": "taxi",
        "amount": "550.00",
        "currency": "ILS",
        "from_location": "Ben Gurion Airport",
        "to_location": "Tel Aviv Hotel",
        "payment_method": "Visa ••1234",
    },
    {
        "id": 2,
        "date": "2026-04-21",
        "time": "12:00",
        "provider": "Yandex Go",
        "category": "taxi",
        "amount": "1200.00",
        "currency": "RUB",
        "from_location": "Sheremetyevo",
        "to_location": "Moscow Center",
        "payment_method": "Mastercard ••5678",
    },
    {
        "id": 3,
        "date": "2026-04-22",
        "time": None,
        "provider": "Careem",
        "category": "taxi",
        "amount": "87.50",
        "currency": "AED",
        "from_location": "DXB Airport",
        "to_location": "Downtown Dubai",
        "payment_method": "Visa ••1234",
    },
    {
        "id": 4,
        "date": "2026-04-22",
        "time": "19:30",
        "provider": "Catch Restaurant",
        "category": "meals",
        "amount": "320.00",
        "currency": "AED",
        "from_location": None,
        "to_location": None,
        "payment_method": "Visa ••1234",
    },
    {
        "id": 5,
        "date": "2026-04-23",
        "time": "09:00",
        "provider": "Uber",
        "category": "taxi",
        "amount": "45.00",
        "currency": "AED",
        "from_location": "Marina",
        "to_location": "DIFC",
        "payment_method": None,
    },
    {
        "id": 6,
        "date": "2026-04-24",
        "time": "14:00",
        "provider": "Bolt",
        "category": "taxi",
        "amount": "62.00",
        "currency": "AED",
        "from_location": "JLT",
        "to_location": "Mall of the Emirates",
        "payment_method": "Apple Pay",
    },
    {
        "id": 7,
        "date": "2026-04-25",
        "time": "10:30",
        "provider": "Marriott",
        "category": "hotel",
        "amount": "850.00",
        "currency": "AED",
        "from_location": None,
        "to_location": None,
        "payment_method": "Visa ••1234",
    },
]


@pytest.fixture
def mock_fx() -> FxRateProvider:
    fx = FxRateProvider()
    rates = {
        ("ils", date(2026, 4, 20)): Decimal("1.01234"),
        ("rub", date(2026, 4, 21)): Decimal("0.04123"),
    }

    async def get_rate(on_date: date, currency: str) -> Decimal:
        c = currency.lower()
        if c == "aed":
            return Decimal("1")
        key = (c, on_date)
        if key in rates:
            return rates[key]
        raise RuntimeError(f"No rate for {currency} on {on_date}")

    fx.get_rate = get_rate  # type: ignore[assignment]
    return fx


@pytest.mark.asyncio
async def test_report_structure(mock_fx: FxRateProvider, tmp_path: Path) -> None:
    out = tmp_path / "test_report.xlsx"
    start = date(2026, 4, 1)
    end = date(2026, 4, 30)

    await build_report(SAMPLE_RECEIPTS, start, end, out, mock_fx)

    assert out.exists()
    wb = load_workbook(out)
    ws = wb.active
    assert ws is not None
    assert ws.title == "Expense Report"

    assert "Expense Report" in str(ws["A1"].value)
    assert "April 2026" in str(ws["A1"].value)

    assert ws.cell(row=4, column=1).value == "#"
    assert ws.cell(row=4, column=2).value == "Date"
    assert ws.cell(row=4, column=8).value == "Amount (AED)"

    assert ws.cell(row=5, column=1).value == 1
    assert ws.cell(row=5, column=2).value == "2026-04-20"
    assert ws.cell(row=5, column=3).value == "GetTaxi"
    assert ws.cell(row=5, column=6).value == "ILS"

    aed_cell = ws.cell(row=5, column=8)
    assert aed_cell.value is not None
    assert "E5*G5" in str(aed_cell.value)

    total_row = 5 + len(SAMPLE_RECEIPTS)
    assert ws.cell(row=total_row, column=4).value == "TOTAL"
    total_formula = str(ws.cell(row=total_row, column=8).value)
    assert "SUM" in total_formula

    wb.close()


@pytest.mark.asyncio
async def test_report_all_formulas_not_hardcoded(
    mock_fx: FxRateProvider, tmp_path: Path
) -> None:
    out = tmp_path / "test_formulas.xlsx"
    await build_report(
        SAMPLE_RECEIPTS, date(2026, 4, 1), date(2026, 4, 30), out, mock_fx
    )

    wb = load_workbook(out)
    ws = wb.active
    assert ws is not None

    for row_idx in range(5, 5 + len(SAMPLE_RECEIPTS)):
        cell = ws.cell(row=row_idx, column=8)
        val = str(cell.value)
        assert val.startswith("="), f"Row {row_idx} col H is hardcoded: {val}"

    wb.close()
